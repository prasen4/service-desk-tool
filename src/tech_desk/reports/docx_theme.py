"""Shared Cotiviti-branded .docx theme.

Colors, floating cover/header artwork, footer/page-number chrome, heading
styles, and table cell styling used by every generated .docx (Position
Papers and Technology Desk Reports) so they all share the exact same
branded look, lifted field-by-field from Cotiviti's official "Position
Paper Template" (.odt).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from tech_desk.timeutils import now_utc

PURPLE = RGBColor(0x30, 0x00, 0x6F)
PINK = RGBColor(0xEC, 0x00, 0x8C)
GRAY = RGBColor(0x5E, 0x61, 0x79)
BORDER_GRAY = "A6A6A6"

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
LOGO_PATH = ASSETS_DIR / "cotiviti_logo.png"
BANNER_PATH = ASSETS_DIR / "cover_banner.png"


def add_floating_picture(
    paragraph, image_path: Path, width_in: float, height_in: float | None, x_in: float, y_in: float, behind: bool = True
) -> None:
    """Insert a page-anchored (floating) picture — used for the cover page's
    full-height banner graphic and logo, matching the template's header
    artwork rather than an inline image that would push text.

    height_in may be None, in which case python-docx computes it from the
    image's own native aspect ratio (avoids stretching/squashing logos that
    don't exactly match a hardcoded width:height ratio).
    """
    if not image_path.exists():
        return
    run = paragraph.add_run()
    if height_in is None:
        run.add_picture(str(image_path), width=Inches(width_in))
    else:
        run.add_picture(str(image_path), width=Inches(width_in), height=Inches(height_in))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    frame_locks = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))

    anchor = OxmlElement("wp:anchor")
    for attr, val in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "251659264",
        "behindDoc": "1" if behind else "0", "locked": "0",
        "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(attr, val)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "page")
    pos_h_off = OxmlElement("wp:posOffset")
    pos_h_off.text = str(int(Inches(x_in)))
    pos_h.append(pos_h_off)
    anchor.append(pos_h)

    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    pos_v_off = OxmlElement("wp:posOffset")
    pos_v_off.text = str(int(Inches(y_in)))
    pos_v.append(pos_v_off)
    anchor.append(pos_v)

    anchor.append(extent)
    effect_extent = OxmlElement("wp:effectExtent")
    for attr in ("l", "t", "r", "b"):
        effect_extent.set(attr, "0")
    anchor.append(effect_extent)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(doc_pr)
    anchor.append(frame_locks)
    anchor.append(graphic)

    drawing.remove(inline)
    drawing.append(anchor)


def add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_end):
        run._r.append(el)
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def border_cell(cell, hex_color: str = BORDER_GRAY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hex_color)
        borders.append(el)
    tc_pr.append(borders)


def set_cell(cell, text: str, *, header: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if center or header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(10)
    if header:
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(cell, PURPLE.__str__())
    border_cell(cell)


def add_heading(doc, text: str, level: int = 1) -> None:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    if level == 1:
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = PURPLE
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = PINK
    else:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_bullets(doc, items, empty_text: str = "No evidence available.") -> None:
    if not items:
        doc.add_paragraph(empty_text)
        return
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def add_body(doc, text: str, empty_text: str = "No evidence available.") -> None:
    doc.add_paragraph(text or empty_text)


def new_branded_document(
    *,
    title: str,
    subtitle: str,
    prepared_for: str = "Prepared for: Cotiviti Technology Leadership",
    delivery_label: str = "For the Delivery of",
    delivery_value: str = "Cotiviti Enterprise AI R&D",
    date_line: str | None = None,
) -> Document:
    """Build a new Document with the Cotiviti cover page, running header/
    footer chrome, and base fonts already applied — matching the official
    Position Paper template's cover layout and body-page master. Callers add
    their own body content (headings/paragraphs/tables) after this returns;
    a page break has already been inserted before the first body page."""
    doc = Document()

    # --- Page geometry (US Letter, matches the template's margins) ---
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    base_font = doc.styles["Normal"].font
    base_font.name = "Calibri"
    base_font.size = Pt(11)

    # --- Cover page has its own header (logo + banner graphic, no footer);
    # body pages get a small header logo + copyright footer — mirrors the
    # template's two master pages (MP0 cover / MP1 body). ---
    section.different_first_page_header_footer = True

    fp_header = section.first_page_header
    fp_header.is_linked_to_previous = False
    fp_header_p = fp_header.paragraphs[0]
    fp_header_p.text = ""
    add_floating_picture(fp_header_p, LOGO_PATH, 2.36, None, 1.0, 0.75, behind=False)
    add_floating_picture(fp_header_p, BANNER_PATH, 2.84, 8.63, 5.66, 0.0, behind=True)

    fp_footer = section.first_page_footer
    fp_footer.is_linked_to_previous = False
    fp_footer.paragraphs[0].text = ""

    header = section.header
    header.is_linked_to_previous = False
    header_p = header.paragraphs[0]
    header_p.text = ""
    add_floating_picture(header_p, LOGO_PATH, 1.53, None, 5.97, 0.3, behind=False)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    footer_p.text = ""
    footer_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    copyright_run = footer_p.add_run(
        f"\u00a9 {now_utc().year} Cotiviti, Inc. All rights reserved. All proprietary information "
        "shall remain the sole and exclusive property of Cotiviti, Inc."
    )
    copyright_run.font.size = Pt(8)
    copyright_run.font.color.rgb = GRAY
    footer_p.add_run("\t")
    add_page_number_field(footer_p)

    # --- Cover page content (mirrors the template's title block, "For the
    # Delivery of" / contact panel, and classification/date lines) ---
    for _ in range(4):
        doc.add_paragraph()

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(36)
    title_run.font.bold = True
    title_run.font.color.rgb = PURPLE

    subtitle_p = doc.add_paragraph()
    subtitle_run = subtitle_p.add_run(subtitle)
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(36)
    subtitle_run.font.color.rgb = PURPLE

    prepared_for_p = doc.add_paragraph()
    prepared_for_run = prepared_for_p.add_run(prepared_for)
    prepared_for_run.font.size = Pt(14)
    prepared_for_run.font.color.rgb = PINK

    for _ in range(3):
        doc.add_paragraph()

    delivery_p = doc.add_paragraph()
    delivery_run = delivery_p.add_run(delivery_label)
    delivery_run.font.size = Pt(12)
    delivery_run.font.color.rgb = PURPLE
    doc.add_paragraph(delivery_value)

    for _ in range(2):
        doc.add_paragraph()

    contact_p = doc.add_paragraph()
    contact_run = contact_p.add_run("Contact")
    contact_run.font.size = Pt(12)
    contact_run.font.color.rgb = PURPLE

    for label in ("Name", "Title", "Cotiviti, Inc.", "Phone"):
        p = doc.add_paragraph()
        p.add_run(label).font.color.rgb = GRAY
    doc.add_paragraph("Email")

    for _ in range(2):
        doc.add_paragraph()
    doc.add_paragraph("Data Sensitivity Classification: Select from drop-down menu")
    doc.add_paragraph(date_line or now_utc().strftime("Prepared %B %d, %Y"))

    doc.add_page_break()
    return doc
